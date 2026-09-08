"""
Step 2: Ingestion.

Reads every file in data/raw/ and turns it into a PolicyDocument: plain text
plus the metadata (doc_id, version, region, effective_date) we need later for
region-aware retrieval and conflict detection.

Supported formats:
- .md / .txt : our HR docs use a small structured header
    ("**Region:** ...", "**Document ID:** ...", etc.) which we parse with a
    regex. This is a stand-in for whatever structured metadata a real HR
    document management system would give us.
- .pdf       : parsed with pypdf. Real PDFs usually don't have this metadata
    in the text, so we fall back to sensible defaults and rely on a filename
    convention instead.
- .docx      : parsed with python-docx, same fallback behavior as PDF.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

from hr_rag.config import PROCESSED_DATA_DIR, RAW_DATA_DIR
from hr_rag.schemas import PolicyDocument

# Matches lines like "**Region:** United States" -> ("Region", "United States")
_METADATA_LINE = re.compile(r"^\*\*(?P<key>[\w ]+):\*\*\s*(?P<value>.+)$", re.MULTILINE)

# Raw "**Region:**" values are free-text descriptions (e.g. "European Union
# (Germany, France, Netherlands offices)"), but region-aware retrieval needs
# to filter on a stable code. Normalize to short codes here, once, at ingest
# time, rather than repeating fuzzy matching logic at query time.
_REGION_CODES = {
    "united states": "US",
    "european union": "EU",
    "global": "Global",
}


def _normalize_region(raw_region: str) -> str:
    lowered = raw_region.lower()
    for keyword, code in _REGION_CODES.items():
        if lowered.startswith(keyword):
            return code
    return raw_region


def _parse_markdown(path: Path) -> PolicyDocument:
    raw = path.read_text(encoding="utf-8")

    fields = {m.group("key").strip(): m.group("value").strip() for m in _METADATA_LINE.finditer(raw)}

    # Title is the first line, e.g. "# US Parental Leave Policy"
    title_match = re.search(r"^#\s+(.+)$", raw, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else path.stem

    # "Document ID" line looks like "HR-POL-US-014 (v2.0) - SUPERSEDES v1.0"
    doc_id_raw = fields.get("Document ID", path.stem)
    doc_id_match = re.match(r"([\w-]+)(?:\s*\((?P<version>v[\d.]+)\))?", doc_id_raw)
    doc_id = doc_id_match.group(1) if doc_id_match else doc_id_raw
    version = doc_id_match.group("version") if doc_id_match and doc_id_match.group("version") else "v1.0"

    return PolicyDocument(
        doc_id=doc_id,
        version=version,
        region=_normalize_region(fields.get("Region", "Unknown")),
        effective_date=fields.get("Effective Date", "Unknown"),
        title=title,
        source_file=path.name,
        text=raw,
    )


def _parse_pdf(path: Path) -> PolicyDocument:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return PolicyDocument(
        doc_id=path.stem,
        version="v1.0",
        region="Unknown",
        effective_date="Unknown",
        title=path.stem,
        source_file=path.name,
        text=text,
    )


def _parse_docx(path: Path) -> PolicyDocument:
    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    return PolicyDocument(
        doc_id=path.stem,
        version="v1.0",
        region="Unknown",
        effective_date="Unknown",
        title=path.stem,
        source_file=path.name,
        text=text,
    )


_PARSERS = {
    ".md": _parse_markdown,
    ".txt": _parse_markdown,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
}


def load_documents(raw_dir: Path = RAW_DATA_DIR) -> list[PolicyDocument]:
    """Parse every supported file in raw_dir into a list of PolicyDocument."""
    documents: list[PolicyDocument] = []
    for path in sorted(raw_dir.iterdir()):
        parser = _PARSERS.get(path.suffix.lower())
        if parser is None:
            continue
        documents.append(parser(path))
    return documents


def save_documents(documents: list[PolicyDocument], out_dir: Path = PROCESSED_DATA_DIR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "documents.json"
    out_path.write_text(
        json.dumps([d.model_dump() for d in documents], indent=2),
        encoding="utf-8",
    )
    return out_path
